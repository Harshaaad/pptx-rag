from pptx import Presentation

def extract_text_from_pptx(file_path: str) -> list[str]:
    prs = Presentation(file_path)
    slides_text = []
    for slide in prs.slides:
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        slides_text.append('\n'.join(slide_text))

    return slides_text[3:]